import os
import logging
from typing import BinaryIO
from langchain_core.documents import Document

logger = logging.getLogger(__name__)


async def load_from_url(url: str) -> list[Document]:
    """Load document content from a URL using Firecrawl."""
    try:
        from backend.mcp.tools.firecrawl_tools import _web_scrape
        result = await _web_scrape({"url": url, "formats": ["markdown"]})
        content = result.get("content", "")
        metadata = result.get("metadata", {})
        metadata["source"] = url
        metadata["loader"] = "firecrawl"
        if not content:
            logger.warning(f"No content scraped from {url}")
            return []
        return [Document(page_content=content, metadata=metadata)]
    except Exception as e:
        logger.error(f"Failed to load from URL {url}: {e}")
        return []


async def load_from_crawl(url: str, max_depth: int = 2, max_pages: int = 10) -> list[Document]:
    """Crawl a website and load all pages as documents using Firecrawl."""
    try:
        from backend.mcp.tools.firecrawl_tools import _web_crawl
        result = await _web_crawl({"url": url, "max_depth": max_depth, "max_pages": max_pages})
        pages = result.get("pages", [])
        docs = []
        for page in pages:
            content = page.get("content", "")
            metadata = page.get("metadata", {})
            metadata["source"] = page.get("url", url)
            metadata["loader"] = "firecrawl_crawl"
            if content:
                docs.append(Document(page_content=content, metadata=metadata))
        return docs
    except Exception as e:
        logger.error(f"Failed to crawl {url}: {e}")
        return []


async def load_document(file: BinaryIO, filename: str) -> list[Document]:
    ext = os.path.splitext(filename)[1].lower()
    docs = []

    try:
        if ext == ".pdf":
            docs = await _load_pdf(file, filename)
        elif ext == ".docx":
            docs = await _load_docx(file, filename)
        elif ext in (".txt", ".md"):
            docs = _load_text(file, filename)
        elif ext == ".csv":
            docs = _load_csv(file, filename)
        elif ext in (".xlsx", ".xls"):
            docs = _load_xlsx(file, filename)
        else:
            logger.warning(f"Unsupported file type: {ext}")
            content = file.read().decode("utf-8", errors="replace")
            docs = [Document(page_content=content, metadata={"filename": filename, "file_type": ext})]
    except Exception as e:
        logger.error(f"Failed to load {filename}: {e}")
        content = file.read().decode("utf-8", errors="replace") if hasattr(file, 'read') else ""
        docs = [Document(page_content=content, metadata={"filename": filename, "file_type": ext, "error": str(e)})]

    for doc in docs:
        doc.metadata["filename"] = filename
        doc.metadata.setdefault("file_type", ext)

    return docs


async def _load_pdf(file: BinaryIO, filename: str) -> list[Document]:
    try:
        from unstructured.partition.auto import partition_pdf
        import tempfile
        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp:
            tmp.write(file.read())
            tmp_path = tmp.name
        elements = partition_pdf(filename=tmp_path)
        os.unlink(tmp_path)
        return [Document(page_content=str(el), metadata={"page": i}) for i, el in enumerate(elements)]
    except ImportError:
        logger.info("unstructured not available, falling back to PyPDFLoader")
        return _load_pdf_fallback(file, filename)


def _load_pdf_fallback(file: BinaryIO, filename: str) -> list[Document]:
    try:
        from langchain_community.document_loaders import PyPDFLoader
        import tempfile
        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp:
            tmp.write(file.read())
            tmp_path = tmp.name
        loader = PyPDFLoader(tmp_path)
        docs = loader.load()
        os.unlink(tmp_path)
        return docs
    except ImportError:
        logger.warning("PyPDFLoader not available, reading raw text")
        content = file.read().decode("utf-8", errors="replace")
        return [Document(page_content=content, metadata={"filename": filename})]


async def _load_docx(file: BinaryIO, filename: str) -> list[Document]:
    try:
        from unstructured.partition.auto import partition_docx
        import tempfile
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
            tmp.write(file.read())
            tmp_path = tmp.name
        elements = partition_docx(filename=tmp_path)
        os.unlink(tmp_path)
        return [Document(page_content=str(el), metadata={"page": i}) for i, el in enumerate(elements)]
    except ImportError:
        content = file.read().decode("utf-8", errors="replace")
        return [Document(page_content=content, metadata={"filename": filename})]


def _load_text(file: BinaryIO, filename: str) -> list[Document]:
    content = file.read().decode("utf-8", errors="replace")
    return [Document(page_content=content, metadata={"filename": filename})]


def _load_csv(file: BinaryIO, filename: str) -> list[Document]:
    import csv
    import io
    text = file.read().decode("utf-8", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    rows = []
    for row in reader:
        rows.append(" | ".join(f"{k}: {v}" for k, v in row.items()))
    return [Document(page_content="\n".join(rows), metadata={"filename": filename, "row_count": len(rows)})]


def _load_xlsx(file: BinaryIO, filename: str) -> list[Document]:
    try:
        import openpyxl
        import io
        wb = openpyxl.load_workbook(io.BytesIO(file.read()), read_only=True)
        docs = []
        for sheet in wb.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                rows.append(" | ".join(str(c) for c in row if c is not None))
            docs.append(Document(
                page_content="\n".join(rows),
                metadata={"filename": filename, "sheet": sheet.title},
            ))
        return docs
    except ImportError:
        content = file.read().decode("utf-8", errors="replace")
        return [Document(page_content=content, metadata={"filename": filename})]
